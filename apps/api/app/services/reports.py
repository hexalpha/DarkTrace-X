from datetime import UTC, datetime
from io import BytesIO
from xml.sax.saxutils import escape
from docx import Document
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


class ReportService:
    def render(self, tenant_id, format_name, request, alerts, iocs=None, cves=None):
        sections = []
        if request.include_alerts:
            sections.append(('Priority alerts', ['Severity', 'Score', 'Status', 'Alert'], [[a.severity.upper(), str(a.score), a.status, a.title] for a in alerts]))
        if request.include_iocs:
            sections.append(('Indicators', ['Type', 'Risk', 'Source', 'Observable'], [[i.type.value, str(i.risk_score), i.source, i.value] for i in iocs or []]))
        if request.include_cves:
            sections.append(('Known exploited vulnerabilities', ['CVE', 'CVSS', 'Product', 'Mitigation'], [[c.cve_id, str(c.cvss) if c.cvss is not None else 'Not supplied', ', '.join(c.affected_products), c.mitigation] for c in cves or []]))
        stamp = f'Workspace: {tenant_id} | Generated: {datetime.now(UTC).isoformat()}'
        if format_name == 'pdf':
            body = self._pdf(request.title, stamp, sections)
            return 'application/pdf', 'darktracex-brief.pdf', body
        if format_name == 'docx':
            body = self._docx(request.title, stamp, sections)
            return 'application/vnd.openxmlformats-officedocument.wordprocessingml.document', 'darktracex-brief.docx', body
        if format_name == 'xlsx':
            body = self._xlsx(request.title, stamp, sections)
            return 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', 'darktracex-brief.xlsx', body
        raise ValueError('Unsupported report format')

    def _pdf(self, title, stamp, sections):
        output = BytesIO()
        styles = getSampleStyleSheet()
        styles['BodyText'].fontSize = 8
        styles['BodyText'].leading = 11
        def para(text):
            return Paragraph(escape(str(text)), styles['BodyText'])
        story = [Paragraph(escape(title), styles['Title']), para(stamp), Spacer(1, 14), para('Source-backed intelligence. Unknown scores are not estimates. Review evidence before taking action.')]
        for heading, headers, rows in sections:
            story.extend([Spacer(1, 14), Paragraph(heading, styles['Heading2'])])
            if not rows:
                story.append(para('No records available.'))
                continue
            table = Table([[para(h) for h in headers], *[[para(v) for v in row] for row in rows]], colWidths=[78, 44, 108, 293], repeatRows=1, hAlign='LEFT')
            table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#DCE8F4')),('GRID',(0,0),(-1,-1),.25,colors.HexColor('#CBD5E1')),('VALIGN',(0,0),(-1,-1),'TOP'),('TOPPADDING',(0,0),(-1,-1),7),('BOTTOMPADDING',(0,0),(-1,-1),7)]))
            story.append(table)
        def footer(canvas, doc):
            canvas.setFont('Helvetica', 8)
            canvas.drawString(36, 20, 'DarkTrace X | Tenant-scoped intelligence')
            canvas.drawRightString(A4[0]-36, 20, f'Page {doc.page}')
        SimpleDocTemplate(output, pagesize=A4, title=title, author='DarkTrace X', leftMargin=36, rightMargin=36, topMargin=36, bottomMargin=36).build(story, onFirstPage=footer, onLaterPages=footer)
        return output.getvalue()

    def _docx(self, title, stamp, sections):
        document=Document(); document.add_heading(title, 0); document.add_paragraph(stamp)
        for heading, headers, rows in sections:
            document.add_heading(heading, 1)
            if not rows:
                document.add_paragraph('No records available.'); continue
            table=document.add_table(rows=1,cols=4); table.style='Light Shading Accent 1'
            for cell, value in zip(table.rows[0].cells,headers): cell.text=value
            for row in rows:
                for cell,value in zip(table.add_row().cells,row): cell.text=str(value)
        output=BytesIO();document.save(output);return output.getvalue()

    def _xlsx(self, title, stamp, sections):
        workbook=Workbook(); summary=workbook.active;summary.title='Summary'
        def safe(value):
            text=str(value)
            return "'"+text if text.startswith(('=','+','-','@','\t','\r')) else text
        summary.append([safe(title)]);summary.append([safe(stamp)])
        for heading,headers,rows in sections:
            sheet=workbook.create_sheet(heading[:31]);sheet.append(headers)
            for row in rows: sheet.append([safe(v) for v in row])
            for cell in sheet[1]: cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='10233D')
            sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
            for column,width in zip('ABCD',[24,16,38,90]): sheet.column_dimensions[column].width=width
        output=BytesIO();workbook.save(output);return output.getvalue()


report_service=ReportService()
